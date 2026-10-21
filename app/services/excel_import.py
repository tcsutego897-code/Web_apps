from __future__ import annotations

import logging
import math
import posixpath
import zipfile
import zlib
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from io import BytesIO
from typing import Any
from xml.etree import ElementTree
from xml.parsers.expat import ExpatError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.config import get_settings
from app.database import get_engine
from app.models import Dataset, DatasetField, ImportHistory, ImportIssue, Record

logger = logging.getLogger(__name__)
PREVIEW_ROW_LIMIT = 5
RECORD_BATCH_SIZE = 500
MAX_ZIP_ENTRIES = 10_000
MAX_MERGED_RANGE_WARNINGS = 5
MAX_CELL_WARNINGS = 10
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


class ImportInputError(ValueError):
    """A user-facing error that should prevent an import."""

    def __init__(self, message: str, source_row_number: int | None = None) -> None:
        super().__init__(message)
        self.source_row_number = source_row_number


@dataclass(frozen=True)
class FieldPreview:
    key: str
    source_name: str
    display_name: str
    position: int


@dataclass(frozen=True)
class PreviewRow:
    row_number: int
    values: list[Any]


@dataclass(frozen=True)
class WorkbookPreview:
    filename: str
    sheet_name: str
    visible_sheets: list[str]
    header_row: int
    fields: list[FieldPreview]
    sample_rows: list[PreviewRow]
    record_count: int
    warnings: list[str]


@dataclass(frozen=True)
class ImportResult:
    dataset_id: int
    dataset_name: str
    record_count: int
    warnings: list[str]


def _validate_archive(content: bytes) -> None:
    settings = get_settings()
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_ENTRIES:
                raise ImportInputError("Excelファイル内のファイル数が上限を超えています。")
            if "[Content_Types].xml" not in archive.namelist() or (
                "xl/workbook.xml" not in archive.namelist()
            ):
                raise ImportInputError("有効な .xlsx ファイルではありません。")
            if any(member.flag_bits & 0x1 for member in members):
                raise ImportInputError("暗号化されたExcelファイルは取り込めません。")
            if any(member.filename.casefold().endswith("vbaproject.bin") for member in members):
                raise ImportInputError("マクロを含むファイルは取り込めません。")
            total_uncompressed_size = sum(member.file_size for member in members)
            if total_uncompressed_size > settings.max_workbook_uncompressed_bytes:
                raise ImportInputError(
                    "展開後のファイルサイズが上限を超えています。"
                )
    except zipfile.BadZipFile as error:
        raise ImportInputError("Excelファイルが壊れているか、.xlsx形式ではありません。") from error
    except zipfile.LargeZipFile as error:
        raise ImportInputError("Excelファイル内のZIP構造を読み取れません。") from error


def _open_workbook_pair(content: bytes) -> tuple[Any, Any]:
    _validate_archive(content)
    formula_workbook = None
    try:
        formula_workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        value_workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (
        InvalidFileException,
        KeyError,
        OSError,
        ValueError,
        IndexError,
        ExpatError,
        zipfile.BadZipFile,
        zlib.error,
    ) as error:
        if formula_workbook is not None:
            formula_workbook.close()
        raise ImportInputError(
            "Excelの内容を読み取れません。ファイルが壊れていないか確認してください。"
        ) from error
    return formula_workbook, value_workbook


def _read_import_row(
    formula_cells: Sequence[Any],
    value_cells: Sequence[Any],
    fields: Sequence[FieldPreview],
    row_number: int,
    cell_warnings: list[str],
) -> dict[str, Any]:
    try:
        return _read_record(
            formula_cells,
            value_cells,
            fields,
            row_number,
            cell_warnings,
        )
    except ImportInputError:
        raise
    except (TypeError, ValueError, OverflowError) as error:
        raise ImportInputError(
            f"Excelの{row_number}行目の値を読み取れません。",
            source_row_number=row_number,
        ) from error


def _build_fields(header_cells: Sequence[Any]) -> tuple[list[FieldPreview], list[str]]:
    fields: list[FieldPreview] = []
    warnings: list[str] = []
    seen_names: set[str] = set()

    for index, cell in enumerate(header_cells):
        raw_value = cell.value
        source_name = "" if raw_value is None else str(raw_value)
        normalized_name = source_name.strip().casefold()
        position = index
        display_name = source_name.strip() or f"列 {index + 1}"
        fields.append(
            FieldPreview(
                key=f"field_{index + 1}",
                source_name=source_name,
                display_name=display_name,
                position=position,
            )
        )
        if not normalized_name:
            warnings.append(f"{index + 1}列目の見出しが空です。表示名を設定してください。")
        elif normalized_name in seen_names:
            warnings.append(
                f"見出し「{source_name}」が重複しています。列は別々に取り込みます。"
            )
        else:
            seen_names.add(normalized_name)

    if not fields or all(not field.source_name.strip() for field in fields):
        raise ImportInputError("見出し行に項目名がありません。見出し行を選び直してください。")
    return fields, warnings


def _merged_ranges(content: bytes, sheet_name: str) -> list[str]:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            relationships_root = ElementTree.fromstring(
                archive.read("xl/_rels/workbook.xml.rels")
            )
            relationships = {
                relation.attrib["Id"]: relation.attrib["Target"]
                for relation in relationships_root.findall(
                    f"{{{PACKAGE_REL_NS}}}Relationship"
                )
            }
            worksheet_path = ""
            for sheet in workbook_root.findall(f".//{{{MAIN_NS}}}sheet"):
                if sheet.attrib.get("name") == sheet_name:
                    target = relationships.get(sheet.attrib.get(f"{{{REL_NS}}}id", ""))
                    if target:
                        if target.startswith("/"):
                            worksheet_path = posixpath.normpath(target.lstrip("/"))
                        elif target.startswith("xl/"):
                            worksheet_path = posixpath.normpath(target)
                        else:
                            worksheet_path = posixpath.normpath(
                                posixpath.join("xl", target)
                            )
                    break
            if not worksheet_path or worksheet_path not in archive.namelist():
                return []

            ranges: list[str] = []
            with archive.open(worksheet_path) as worksheet_xml:
                for _, element in ElementTree.iterparse(
                    worksheet_xml,
                    events=("end",),
                ):
                    if element.tag == f"{{{MAIN_NS}}}mergeCell":
                        if len(ranges) < MAX_MERGED_RANGE_WARNINGS:
                            ranges.append(element.attrib.get("ref", ""))
                    element.clear()
            return ranges
    except (zipfile.BadZipFile, KeyError, OSError, ElementTree.ParseError) as error:
        raise ImportInputError("結合セル情報を読み取れません。") from error


def _normalize_cell_value(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _read_record(
    formula_cells: Sequence[Any],
    value_cells: Sequence[Any],
    fields: Sequence[FieldPreview],
    row_number: int,
    cell_warnings: list[str],
) -> dict[str, Any]:
    data: dict[str, Any] = {}
    for field, formula_cell, value_cell in zip(fields, formula_cells, value_cells):
        if value_cell.data_type == "e":
            raise ImportInputError(
                f"Excelの{row_number}行目（{formula_cell.coordinate}）にエラー値があります。",
                source_row_number=row_number,
            )
        if formula_cell.data_type == "f" and value_cell.value is None:
            if len(cell_warnings) < MAX_CELL_WARNINGS:
                cell_warnings.append(
                    f"{formula_cell.coordinate}: 保存済みの数式結果がないため空欄にしました。"
                )
            value = None
        else:
            value = _normalize_cell_value(value_cell.value)
        data[field.key] = value

    if all(value is None for value in data.values()):
        return {}
    return data


def _read_header_row(sheet: Any, header_row: int, column_count: int) -> list[Any]:
    try:
        return next(
            sheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                max_col=column_count,
            )
        )
    except StopIteration:
        raise ImportInputError("指定した行に見出しがありません。") from None
    except (
        InvalidFileException,
        KeyError,
        OSError,
        ValueError,
        IndexError,
        ExpatError,
        zipfile.BadZipFile,
        zlib.error,
    ) as error:
        raise ImportInputError("Excelの見出し行を読み取れません。") from error


def _inspect_sheet(
    filename: str,
    content: bytes,
    sheet_name: str | None,
    header_row: int,
) -> WorkbookPreview:
    settings = get_settings()
    if header_row < 1:
        raise ImportInputError("見出し行は1以上の行番号を指定してください。")

    formula_workbook, value_workbook = _open_workbook_pair(content)
    try:
        visible_sheets = [
            worksheet.title
            for worksheet in formula_workbook.worksheets
            if worksheet.sheet_state == "visible"
        ]
        if not visible_sheets:
            raise ImportInputError("表示状態のシートがありません。")
        selected_sheet = sheet_name or visible_sheets[0]
        if selected_sheet not in visible_sheets:
            raise ImportInputError("選択したシートが見つからないか、非表示です。")

        formula_sheet = formula_workbook[selected_sheet]
        value_sheet = value_workbook[selected_sheet]
        column_count = formula_sheet.max_column or 0
        if column_count < 1:
            raise ImportInputError("シートに見出しがありません。")
        if header_row > (formula_sheet.max_row or 0):
            raise ImportInputError("見出し行がシートの範囲外です。")

        fields, field_warnings = _build_fields(
            _read_header_row(formula_sheet, header_row, column_count)
        )
        warnings = list(field_warnings)
        cell_warnings: list[str] = []
        sample_rows: list[PreviewRow] = []
        record_count = 0

        formula_rows = formula_sheet.iter_rows(
            min_row=header_row + 1,
            max_col=len(fields),
        )
        value_rows = value_sheet.iter_rows(
            min_row=header_row + 1,
            max_col=len(fields),
        )
        try:
            for row_number, (formula_cells, value_cells) in enumerate(
                zip(formula_rows, value_rows),
                start=header_row + 1,
            ):
                data = _read_import_row(
                    formula_cells,
                    value_cells,
                    fields,
                    row_number,
                    cell_warnings,
                )
                if not data:
                    continue
                record_count += 1
                if record_count > settings.max_import_rows:
                    raise ImportInputError(
                        f"Excelの{row_number}行目で、取り込み可能な"
                        f"データ行数の上限（{settings.max_import_rows:,}行）を超えました。",
                        source_row_number=row_number,
                    )
                if len(sample_rows) < PREVIEW_ROW_LIMIT:
                    sample_rows.append(
                        PreviewRow(
                            row_number=row_number,
                            values=[data[field.key] for field in fields],
                        )
                    )
        except ImportInputError:
            raise
        except (
            InvalidFileException,
            KeyError,
            OSError,
            ValueError,
            IndexError,
            ExpatError,
            zipfile.BadZipFile,
            zlib.error,
        ) as error:
            raise ImportInputError(
                "Excelのデータ行を読み取れません。ファイルが壊れていないか確認してください。"
            ) from error

        if record_count == 0:
            raise ImportInputError(
                "見出しより下にデータがありません。見出し行を確認してください。"
            )

        if cell_warnings:
            warnings.append("数式の計算結果が保存されていないセルがあります。")
            warnings.extend(cell_warnings)
            if len(cell_warnings) == MAX_CELL_WARNINGS:
                warnings.append("表示できる数式の警告は先頭10件までです。")

        ranges = _merged_ranges(content, selected_sheet)
        if ranges:
            displayed_ranges = "、".join(ranges)
            warnings.append(
                f"結合セルがあります（{displayed_ranges}）。"
                "左上セルの値だけを使い、他のセルへ補完しません。"
            )

        return WorkbookPreview(
            filename=filename,
            sheet_name=selected_sheet,
            visible_sheets=visible_sheets,
            header_row=header_row,
            fields=fields,
            sample_rows=sample_rows,
            record_count=record_count,
            warnings=warnings,
        )
    finally:
        formula_workbook.close()
        value_workbook.close()


def preview_workbook(
    filename: str,
    content: bytes,
    sheet_name: str | None,
    header_row: int,
) -> WorkbookPreview:
    return _inspect_sheet(filename, content, sheet_name, header_row)


def list_visible_sheets(content: bytes) -> list[str]:
    formula_workbook, value_workbook = _open_workbook_pair(content)
    try:
        return [
            worksheet.title
            for worksheet in formula_workbook.worksheets
            if worksheet.sheet_state == "visible"
        ]
    finally:
        formula_workbook.close()
        value_workbook.close()


def _iter_import_records(
    content: bytes,
    preview: WorkbookPreview,
) -> Iterator[tuple[int, dict[str, Any]]]:
    formula_workbook, value_workbook = _open_workbook_pair(content)
    try:
        formula_sheet = formula_workbook[preview.sheet_name]
        value_sheet = value_workbook[preview.sheet_name]
        formula_rows = formula_sheet.iter_rows(
            min_row=preview.header_row + 1,
            max_col=len(preview.fields),
        )
        value_rows = value_sheet.iter_rows(
            min_row=preview.header_row + 1,
            max_col=len(preview.fields),
        )
        cell_warnings: list[str] = []
        count = 0
        try:
            for row_number, (formula_cells, value_cells) in enumerate(
                zip(formula_rows, value_rows),
                start=preview.header_row + 1,
            ):
                data = _read_import_row(
                    formula_cells,
                    value_cells,
                    preview.fields,
                    row_number,
                    cell_warnings,
                )
                if data:
                    count += 1
                    if count > get_settings().max_import_rows:
                        raise ImportInputError(
                            f"Excelの{row_number}行目で、プレビュー後に"
                            "データ行数の上限を超えました。",
                            source_row_number=row_number,
                        )
                    yield row_number, data
        except ImportInputError:
            raise
        except (
            InvalidFileException,
            KeyError,
            OSError,
            ValueError,
            IndexError,
            ExpatError,
            zipfile.BadZipFile,
            zlib.error,
        ) as error:
            raise ImportInputError(
                "Excelのデータ行を読み取れません。ファイルが壊れていないか確認してください。"
            ) from error
    finally:
        formula_workbook.close()
        value_workbook.close()


def _clean_display_names(
    fields: Sequence[FieldPreview],
    display_names: Sequence[str] | None,
) -> list[str]:
    if display_names is None:
        return [field.display_name for field in fields]
    if len(display_names) != len(fields):
        raise ImportInputError("項目名の数が列数と一致しません。")

    cleaned_names = []
    for field, display_name in zip(fields, display_names):
        name = display_name.strip() or field.display_name
        if len(name) > 255:
            raise ImportInputError("項目名は255文字以内にしてください。")
        cleaned_names.append(name)
    return cleaned_names


def _finish_history(
    session: Session,
    history_id: int,
    status: str,
    success_count: int,
    failure_count: int,
    error_summary: str | None,
    issue: ImportInputError | None = None,
) -> None:
    history = session.scalar(
        select(ImportHistory).where(ImportHistory.id == history_id)
    )
    if history is None:
        raise RuntimeError("Import history disappeared during processing.")
    history.status = status
    history.success_count = success_count
    history.failure_count = failure_count
    history.error_summary = error_summary
    history.completed_at = datetime.now().astimezone()
    if issue is not None:
        history.issues.append(
            ImportIssue(
                source_row_number=issue.source_row_number,
                message=str(issue),
            )
        )
    session.commit()


def record_failed_preview(
    filename: str,
    sheet_name: str,
    error: ImportInputError,
) -> None:
    with Session(get_engine()) as session:
        history = ImportHistory(
            source_filename=filename,
            sheet_name=sheet_name,
            status="failed",
            success_count=0,
            failure_count=1,
            error_summary=str(error),
            completed_at=datetime.now().astimezone(),
        )
        if error.source_row_number is not None:
            history.issues.append(
                ImportIssue(
                    source_row_number=error.source_row_number,
                    message=str(error),
                )
            )
        session.add(history)
        session.commit()


def import_workbook(
    filename: str,
    content: bytes,
    sheet_name: str,
    header_row: int,
    dataset_name: str,
    display_names: Sequence[str] | None,
) -> ImportResult:
    cleaned_dataset_name = dataset_name.strip()

    engine = get_engine()
    try:
        with Session(engine) as session:
            history = ImportHistory(
                source_filename=filename,
                sheet_name=sheet_name,
                status="pending",
            )
            session.add(history)
            session.flush()
            history_id = history.id
            session.commit()

            try:
                if not cleaned_dataset_name:
                    raise ImportInputError("データセット名を入力してください。")
                if len(cleaned_dataset_name) > 255:
                    raise ImportInputError("データセット名は255文字以内にしてください。")
                preview = _inspect_sheet(filename, content, sheet_name, header_row)
                names = _clean_display_names(preview.fields, display_names)
                dataset = Dataset(
                    name=cleaned_dataset_name,
                    source_filename=filename,
                    sheet_name=sheet_name,
                    record_count=preview.record_count,
                )
                with session.begin():
                    session.add(dataset)
                    session.flush()
                    session.add_all(
                        [
                            DatasetField(
                                dataset_id=dataset.id,
                                key=field.key,
                                source_name=field.source_name,
                                display_name=display_name,
                                position=field.position,
                                is_searchable=True,
                            )
                            for field, display_name in zip(preview.fields, names)
                        ]
                    )

                    record_batch: list[Record] = []
                    imported_count = 0
                    for row_number, data in _iter_import_records(content, preview):
                        record_batch.append(
                            Record(
                                dataset_id=dataset.id,
                                source_row_number=row_number,
                                data=data,
                            )
                        )
                        imported_count += 1
                        if len(record_batch) >= RECORD_BATCH_SIZE:
                            session.add_all(record_batch)
                            session.flush()
                            record_batch.clear()
                    if record_batch:
                        session.add_all(record_batch)
                    if imported_count != preview.record_count:
                        raise ImportInputError(
                            "プレビュー後にデータ行数が変わりました。再度プレビューしてください。"
                        )

                    stored_history = session.get(ImportHistory, history_id)
                    if stored_history is None:
                        raise RuntimeError("Import history disappeared during processing.")
                    stored_history.dataset = dataset
                    stored_history.status = "succeeded"
                    stored_history.success_count = imported_count
                    stored_history.failure_count = 0
                    stored_history.error_summary = None
                    stored_history.completed_at = datetime.now().astimezone()

                return ImportResult(
                    dataset_id=dataset.id,
                    dataset_name=cleaned_dataset_name,
                    record_count=preview.record_count,
                    warnings=preview.warnings,
                )
            except ImportInputError as error:
                session.rollback()
                _finish_history(
                    session,
                    history_id,
                    status="failed",
                    success_count=0,
                    failure_count=1,
                    error_summary=str(error),
                    issue=error if error.source_row_number is not None else None,
                )
                raise
            except SQLAlchemyError:
                session.rollback()
                _finish_history(
                    session,
                    history_id,
                    status="failed",
                    success_count=0,
                    failure_count=1,
                    error_summary="保存処理に失敗したため、取り込み全体を取り消しました。",
                )
                raise
    except SQLAlchemyError:
        logger.exception("Workbook import failed: %s", filename)
        raise
