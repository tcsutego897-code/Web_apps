import asyncio
import re
from datetime import date
from io import BytesIO
from pathlib import Path
from sqlite3 import Connection
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree

import httpx
import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.config import Settings
from app.database import Base
from app.main import app
from app.models import Dataset, DatasetField, ImportHistory, ImportIssue, Record
from app.services.excel_import import (
    ImportInputError,
    import_workbook,
    preview_workbook,
)
from app.services.upload_storage import (
    UploadError,
    read_staged_upload,
    read_upload,
)
from starlette.datastructures import UploadFile


def create_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "People"
    sheet.append(["ID", "Name", "Name", "Birthday", "Formula"])
    sheet.append(["000123", "Alice", "A.", date(2000, 1, 2), "=1+1"])
    sheet.append([None, None, None, None, None])
    sheet.append(["000456", "Bob", "B.", None, None])

    orders_sheet = workbook.create_sheet("Orders")
    orders_sheet.append(["Order ID", "Total"])
    orders_sheet.append([18, 125.5])

    hidden_sheet = workbook.create_sheet("Private")
    hidden_sheet.sheet_state = "hidden"
    hidden_sheet.append(["Secret"])
    hidden_sheet.append(["hidden value"])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_error_workbook() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Code", "Result"])
    worksheet.append(["A-1", "#VALUE!"])
    worksheet["B2"].data_type = "e"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def add_cached_formula_result(content: bytes, cell_reference: str) -> bytes:
    output = BytesIO()
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(BytesIO(content)) as source, ZipFile(
        output,
        mode="w",
        compression=ZIP_DEFLATED,
    ) as destination:
        for item in source.infolist():
            file_content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(file_content)
                for cell in root.iter(f"{namespace}c"):
                    if cell.attrib.get("r") == cell_reference:
                        value = cell.find(f"{namespace}v")
                        assert value is not None
                        value.text = "42"
                file_content = ElementTree.tostring(
                    root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            destination.writestr(item, file_content)
    return output.getvalue()


def add_archive_member(content: bytes, filename: str, member_content: bytes) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(
        output,
        mode="w",
        compression=ZIP_DEFLATED,
    ) as destination:
        for item in source.infolist():
            destination.writestr(item, source.read(item.filename))
        destination.writestr(filename, member_content)
    return output.getvalue()


async def request(
    method: str,
    path: str,
    *,
    data: dict[str, str] | None = None,
    files: dict[str, tuple[str, bytes, str]] | None = None,
) -> httpx.Response:
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(
        transport=transport,
        base_url="http://testserver",
    ) as client:
        return await client.request(method, path, data=data, files=files)


def post(
    path: str,
    *,
    data: dict[str, str] | None = None,
    files: dict[str, tuple[str, bytes, str]] | None = None,
) -> httpx.Response:
    return asyncio.run(request("POST", path, data=data, files=files))


def test_preview_preserves_duplicate_columns_dates_rows_and_warnings() -> None:
    preview = preview_workbook(
        "people.xlsx",
        create_workbook(),
        sheet_name=None,
        header_row=1,
    )

    assert preview.visible_sheets == ["People", "Orders"]
    assert [field.key for field in preview.fields] == [
        "field_1",
        "field_2",
        "field_3",
        "field_4",
        "field_5",
    ]
    assert [field.source_name for field in preview.fields] == [
        "ID",
        "Name",
        "Name",
        "Birthday",
        "Formula",
    ]
    assert preview.record_count == 2
    assert [row.row_number for row in preview.sample_rows] == [2, 4]
    assert preview.sample_rows[0].values == [
        "000123",
        "Alice",
        "A.",
        "2000-01-02T00:00:00",
        None,
    ]
    assert any("重複" in warning for warning in preview.warnings)
    assert any("数式" in warning for warning in preview.warnings)
    assert any("E2" in warning for warning in preview.warnings)


def test_preview_rejects_hidden_sheets_and_invalid_header_rows() -> None:
    content = create_workbook()

    with pytest.raises(ImportInputError, match="非表示"):
        preview_workbook("people.xlsx", content, "Private", 1)
    with pytest.raises(ImportInputError, match="見出し行"):
        preview_workbook("people.xlsx", content, "People", 0)


def test_excel_cell_error_reports_source_row_number() -> None:
    with pytest.raises(ImportInputError, match="2行目") as error:
        preview_workbook(
            "error.xlsx",
            create_error_workbook(),
            sheet_name=None,
            header_row=1,
        )

    assert error.value.source_row_number == 2


def test_each_visible_sheet_uses_its_own_column_definitions() -> None:
    preview = preview_workbook(
        "people.xlsx",
        create_workbook(),
        "Orders",
        1,
    )

    assert [field.source_name for field in preview.fields] == ["Order ID", "Total"]
    assert preview.sample_rows[0].values == [18, 125.5]


def test_preview_supports_a_header_below_the_first_row() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["A title above the table"])
    worksheet.append(["Code", "Amount"])
    worksheet.append(["0007", 42])
    output = BytesIO()
    workbook.save(output)

    preview = preview_workbook("titled.xlsx", output.getvalue(), None, 2)

    assert [field.source_name for field in preview.fields] == ["Code", "Amount"]
    assert preview.sample_rows[0].row_number == 3
    assert preview.sample_rows[0].values == ["0007", 42]


def test_preview_reports_merged_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Merged"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Name"
    sheet.append(["Alice", None])
    output = BytesIO()
    workbook.save(output)

    preview = preview_workbook("merged.xlsx", output.getvalue(), "Merged", 1)

    assert any("A1:B1" in warning for warning in preview.warnings)
    assert preview.sample_rows[0].values == ["Alice", None]


def test_preview_uses_saved_formula_results_without_calculating_formulas() -> None:
    content = add_cached_formula_result(create_workbook(), "E2")

    preview = preview_workbook("people.xlsx", content, "People", 1)

    assert preview.sample_rows[0].values[-1] == 42
    assert not any("数式の計算結果が保存されていない" in warning for warning in preview.warnings)


def test_preview_rejects_bad_file_and_empty_data() -> None:
    with pytest.raises(ImportInputError, match="壊れている"):
        preview_workbook("broken.xlsx", b"not an xlsx file", None, 1)

    workbook = Workbook()
    workbook.active.append(["Name"])
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ImportInputError, match="データがありません"):
        preview_workbook("empty.xlsx", output.getvalue(), None, 1)


def test_preview_rejects_macro_payload_and_oversized_archives() -> None:
    workbook_content = create_workbook()
    macro_workbook = add_archive_member(
        workbook_content,
        "xl/vbaProject.bin",
        b"macro data",
    )
    with pytest.raises(ImportInputError, match="マクロ"):
        preview_workbook("macro.xlsx", macro_workbook, None, 1)

    with patch(
        "app.services.excel_import.get_settings",
        return_value=Settings(max_workbook_uncompressed_bytes=100),
    ):
        with pytest.raises(ImportInputError, match="展開後"):
            preview_workbook("large.xlsx", workbook_content, None, 1)


def test_preview_warns_for_blank_headers_and_enforces_row_limit() -> None:
    workbook = Workbook()
    workbook.active.append(["Name", None])
    workbook.active.append(["Alice", None])
    output = BytesIO()
    workbook.save(output)

    preview = preview_workbook("blank-header.xlsx", output.getvalue(), None, 1)
    assert preview.fields[1].key == "field_2"
    assert preview.fields[1].display_name == "列 2"
    assert any("見出しが空" in warning for warning in preview.warnings)

    with patch(
        "app.services.excel_import.get_settings",
        return_value=Settings(max_import_rows=1),
    ):
        with pytest.raises(ImportInputError, match="行数"):
            preview_workbook("people.xlsx", create_workbook(), "People", 1)


def test_file_upload_checks_format_and_size() -> None:
    from app.services.upload_storage import UploadError

    settings = Settings(max_upload_size_bytes=4)
    with patch("app.services.upload_storage.get_settings", return_value=settings):
        with pytest.raises(UploadError, match="上限"):
            asyncio.run(
                read_upload(
                    UploadFile(
                        filename="too-large.xlsx",
                        file=BytesIO(b"12345"),
                    )
                )
            )

    with pytest.raises(UploadError, match="\\.xlsx"):
        asyncio.run(
            read_upload(
                UploadFile(filename="old.xls", file=BytesIO(b"data"))
            )
        )


def test_preview_and_confirm_import_store_records_and_history(tmp_path: Path) -> None:
    database_path = tmp_path / "imports.db"
    engine = create_engine(f"sqlite+pysqlite:///{database_path}")

    @event.listens_for(engine, "connect")
    def enable_foreign_keys(connection: Connection, _record: object) -> None:
        del _record
        connection.execute("PRAGMA foreign_keys = ON")

    Base.metadata.create_all(engine)
    upload = {
        "upload": (
            "people.xlsx",
            create_workbook(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    try:
        with (
            patch("app.services.excel_import.get_engine", return_value=engine),
            patch("app.routers.pages.get_engine", return_value=engine),
            patch(
                "app.services.upload_storage.STAGING_DIRECTORY",
                tmp_path / "staged-uploads",
            ),
        ):
            preview_response = post(
                "/import/preview",
                data={"header_row": "1"},
                files=upload,
            )
            assert preview_response.status_code == 200
            assert "取り込みプレビュー" in preview_response.text
            assert "Name" in preview_response.text
            assert "2000-01-02" in preview_response.text

            token_match = re.search(
                r'name="staging_token" value="([A-Za-z0-9_-]{43})"',
                preview_response.text,
            )
            assert token_match is not None
            token = token_match.group(1)

            result_response = post(
                "/import/confirm",
                data={
                    "staging_token": token,
                    "filename": "people.xlsx",
                    "sheet_name": "People",
                    "header_row": "1",
                    "dataset_name": "People list",
                    "display_names": [
                        "ID",
                        "Full name",
                        "Alias",
                        "Birthday",
                        "Formula",
                    ],
                },
            )
            assert result_response.status_code == 200
            assert "取り込みが完了しました" in result_response.text
            assert "People list" in result_response.text

            dataset_response = asyncio.run(request("GET", "/"))
            assert dataset_response.status_code == 200
            assert "People list" in dataset_response.text

        with Session(engine) as session:
            dataset = session.scalar(
                select(Dataset).where(Dataset.name == "People list")
            )
            assert dataset is not None
            assert dataset.record_count == 2
            assert dataset.source_filename == "people.xlsx"
            assert dataset.sheet_name == "People"
            assert [field.display_name for field in dataset.fields] == [
                "ID",
                "Full name",
                "Alias",
                "Birthday",
                "Formula",
            ]
            records = session.scalars(
                select(Record)
                .where(Record.dataset_id == dataset.id)
                .order_by(Record.source_row_number)
            ).all()
            assert [record.source_row_number for record in records] == [2, 4]
            assert records[0].data == {
                "field_1": "000123",
                "field_2": "Alice",
                "field_3": "A.",
                "field_4": "2000-01-02T00:00:00",
                "field_5": None,
            }
            assert session.scalar(
                select(ImportHistory).where(ImportHistory.status == "succeeded")
            ) is not None
            with pytest.raises(UploadError, match="ありません"):
                read_staged_upload(token)
    finally:
        engine.dispose()


def test_record_save_failure_rolls_back_dataset_and_records(tmp_path: Path) -> None:
    engine = create_engine(f"sqlite+pysqlite:///{tmp_path / 'rollback.db'}")
    Base.metadata.create_all(engine)

    @event.listens_for(engine, "before_cursor_execute")
    def fail_record_insert(
        _connection: object,
        _cursor: object,
        statement: str,
        _parameters: object,
        _context: object,
        _executemany: bool,
    ) -> None:
        del _connection, _cursor, _parameters, _context, _executemany
        if statement.lstrip().startswith("INSERT INTO records"):
            raise SQLAlchemyError("Simulated record save failure")

    try:
        with patch("app.services.excel_import.get_engine", return_value=engine):
            with pytest.raises(SQLAlchemyError, match="Simulated"):
                import_workbook(
                    "people.xlsx",
                    create_workbook(),
                    "People",
                    1,
                    "Must be rolled back",
                    None,
                )

        with Session(engine) as session:
            assert session.scalars(select(Dataset)).all() == []
            assert session.scalars(select(DatasetField)).all() == []
            assert session.scalars(select(Record)).all() == []
            history = session.scalar(
                select(ImportHistory).where(ImportHistory.status == "failed")
            )
            assert history is not None
            assert history.failure_count == 1
            assert "取り込み全体を取り消しました" in history.error_summary
    finally:
        engine.dispose()


def test_import_validation_failure_is_recorded_and_does_not_create_dataset(
    tmp_path: Path,
) -> None:
    engine = create_engine(f"sqlite+pysqlite:///{tmp_path / 'failed-import.db'}")
    Base.metadata.create_all(engine)
    try:
        with patch("app.services.excel_import.get_engine", return_value=engine):
            with pytest.raises(ImportInputError, match="非表示"):
                import_workbook(
                    "people.xlsx",
                    create_workbook(),
                    "Private",
                    1,
                    "Should not exist",
                    None,
                )

        with Session(engine) as session:
            assert session.scalars(select(Dataset)).all() == []
            failure = session.scalar(
                select(ImportHistory).where(ImportHistory.status == "failed")
            )
            assert failure is not None
            assert failure.failure_count == 1
            assert "非表示" in failure.error_summary
    finally:
        engine.dispose()


def test_preview_row_error_is_saved_to_import_history(tmp_path: Path) -> None:
    engine = create_engine(f"sqlite+pysqlite:///{tmp_path / 'row-error.db'}")
    Base.metadata.create_all(engine)
    try:
        with (
            patch("app.services.excel_import.get_engine", return_value=engine),
            patch(
                "app.services.upload_storage.STAGING_DIRECTORY",
                tmp_path / "row-error-staging",
            ),
        ):
            response = post(
                "/import/preview",
                data={"header_row": "1"},
                files={
                    "upload": (
                        "error.xlsx",
                        create_error_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )

        assert response.status_code == 400
        assert "2行目" in response.text
        with Session(engine) as session:
            history = session.scalar(
                select(ImportHistory).where(ImportHistory.status == "failed")
            )
            assert history is not None
            assert history.failure_count == 1
            assert history.dataset_id is None
            issue = session.scalar(
                select(ImportIssue).where(
                    ImportIssue.import_history_id == history.id
                )
            )
            assert issue is not None
            assert issue.source_row_number == 2
            assert "エラー値" in issue.message
    finally:
        engine.dispose()
